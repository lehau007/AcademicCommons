"""Deterministic PPTX slide → Markdown (python-pptx) with slots for picture shapes."""

from __future__ import annotations

import io
import logging
from collections.abc import Iterator
from dataclasses import dataclass, field
from typing import Any

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

logger = logging.getLogger(__name__)

SKIP_PLACEHOLDERS = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}
TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
EMU_PER_INCH = 914400
MIN_PICTURE_SIDE_EMU = int(0.4 * EMU_PER_INCH)


@dataclass
class SlidePicture:
    slot: str
    index: int  # 1-based per slide
    blob: bytes
    ext: str
    too_small: bool


@dataclass
class SlideMarkdown:
    number: int
    markdown: str
    text: str
    pictures: list[SlidePicture] = field(default_factory=list)


def _clean(text: str) -> str:
    return " ".join(str(text).replace("\x0b", " ").split())


def _shape_type(shape: Any) -> Any:
    """python-pptx raises NotImplementedError for a ``p:sp`` without geometry; treat it as untyped."""
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def _iter_shapes(shapes: Any) -> Iterator[Any]:
    for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if _shape_type(shape) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _placeholder_type(shape: Any) -> Any:
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape.placeholder_format.type
    except ValueError:
        return None


def _is_picture(shape: Any) -> bool:
    if _shape_type(shape) == MSO_SHAPE_TYPE.PICTURE:
        return True
    return bool(getattr(shape, "is_placeholder", False) and getattr(shape, "image", None) is not None)


def text_frame_markdown(frame: Any) -> str:
    items = [(int(p.level or 0), _clean(p.text)) for p in frame.paragraphs]
    items = [(level, text) for level, text in items if text]
    if not items:
        return ""
    if len(items) == 1:
        return items[0][1]
    return "\n".join(f"{'  ' * level}- {text}" for level, text in items)


def table_markdown(table: Any) -> str:
    rows = [[_clean(cell.text).replace("|", "\\|") for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def slide_to_markdown(slide: Any, number: int) -> SlideMarkdown:
    title_shape = slide.shapes.title
    has_title_text = title_shape is not None and getattr(title_shape, "has_text_frame", False)
    title = _clean(title_shape.text) if has_title_text else ""
    parts = [f"## {title}" if title else f"## Slide {number}"]
    texts = [title] if title else []
    pictures: list[SlidePicture] = []
    for shape in _iter_shapes(slide.shapes):
        try:
            kind = _placeholder_type(shape)
            if kind in SKIP_PLACEHOLDERS or kind in TITLE_PLACEHOLDERS:
                continue
            if _is_picture(shape):
                image = shape.image  # ValueError for a linked (not embedded) picture
                index = len(pictures) + 1
                slot = f"@@FIGURE_{number}_{index}@@"
                too_small = min(int(shape.width or 0), int(shape.height or 0)) < MIN_PICTURE_SIDE_EMU
                pictures.append(SlidePicture(slot, index, image.blob, str(image.ext), too_small))
                parts.append(slot)
                continue
            if getattr(shape, "has_table", False):
                markdown = table_markdown(shape.table)
            elif getattr(shape, "has_text_frame", False):
                markdown = text_frame_markdown(shape.text_frame)
            else:
                markdown = ""
        except Exception:  # noqa: BLE001 - one unreadable shape must not fail the whole document
            logger.warning("skipping unreadable shape on slide %s", number, exc_info=True)
            continue
        if markdown:
            parts.append(markdown)
            texts.append(markdown)
    return SlideMarkdown(number=number, markdown="\n\n".join(parts), text="\n".join(texts), pictures=pictures)


def normalize_picture(blob: bytes, ext: str) -> tuple[bytes, str, str] | None:
    """PNG/JPEG pass through; other raster formats are converted to PNG; unreadable → None."""
    ext = ext.lower()
    if ext == "png":
        return blob, "png", "image/png"
    if ext in ("jpg", "jpeg"):
        return blob, "jpg", "image/jpeg"
    try:
        with Image.open(io.BytesIO(blob)) as image:
            mode = "RGBA" if image.mode in ("RGBA", "LA", "P") else "RGB"
            buffer = io.BytesIO()
            image.convert(mode).save(buffer, "PNG")
        return buffer.getvalue(), "png", "image/png"
    except Exception:  # noqa: BLE001 - EMF/WMF and corrupt blobs are skipped
        return None
