from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches

from app.services.document_processing.pptx_markdown import normalize_picture, slide_to_markdown

from .conftest import png_bytes


def _deck(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Graph basics"
    body = slide.placeholders[1].text_frame
    body.text = "Vertices"
    para = body.add_paragraph()
    para.text = "Edges connect | vertices"
    para.level = 1
    slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(5), Inches(4), Inches(3), Inches(2))
    slide.shapes.add_picture(io.BytesIO(png_bytes(60, 40)), Inches(0.1), Inches(0.1), Inches(0.3), Inches(0.3))
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(5.5), Inches(4), Inches(1)).table
    for (r, c), value in {(0, 0): "A", (0, 1): "B", (1, 0): "1", (1, 1): "2"}.items():
        table.cell(r, c).text = value
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def test_slide_markdown_title_bullets_table_and_picture_slots(tmp_path: Path) -> None:
    slide = Presentation(str(_deck(tmp_path))).slides[0]

    result = slide_to_markdown(slide, 1)

    assert result.markdown.startswith("## Graph basics")
    assert "- Vertices\n  - Edges connect | vertices" in result.markdown
    assert "| A | B |\n| --- | --- |\n| 1 | 2 |" in result.markdown
    assert [p.slot for p in result.pictures] == ["@@FIGURE_1_1@@", "@@FIGURE_1_2@@"]
    assert [p.too_small for p in result.pictures] == [True, False]  # sorted top-to-bottom: logo first
    assert all(p.slot in result.markdown for p in result.pictures)
    assert "Vertices" in result.text


def test_footer_date_and_slide_number_placeholders_are_skipped() -> None:
    def _ph(kind: object, text: str) -> SimpleNamespace:
        frame = SimpleNamespace(paragraphs=[SimpleNamespace(level=0, text=text)])
        return SimpleNamespace(
            shape_type=MSO_SHAPE_TYPE.PLACEHOLDER, is_placeholder=True, top=0, left=0,
            placeholder_format=SimpleNamespace(type=kind), has_text_frame=True, text_frame=frame, has_table=False,
        )

    shapes = [_ph(PP_PLACEHOLDER.FOOTER, "Bộ môn KHMT"), _ph(PP_PLACEHOLDER.SLIDE_NUMBER, "7"),
              _ph(PP_PLACEHOLDER.BODY, "Real content")]
    slide = SimpleNamespace(shapes=_Shapes(shapes))

    result = slide_to_markdown(slide, 7)

    assert result.markdown == "## Slide 7\n\nReal content"


class _Shapes(list):  # type: ignore[type-arg]
    """Iterable shape collection with python-pptx's ``.title`` attribute."""

    title = None


class _LinkedPicture:
    """A picture whose image is linked, not embedded: python-pptx raises ValueError on ``.image``."""

    shape_type = MSO_SHAPE_TYPE.PICTURE
    is_placeholder = False
    top, left, width, height = Inches(1), Inches(1), Inches(3), Inches(2)

    @property
    def image(self) -> object:
        raise ValueError("no embedded image")


class _LinkedPlaceholderPicture(_LinkedPicture):
    shape_type = MSO_SHAPE_TYPE.PLACEHOLDER
    is_placeholder = True
    placeholder_format = SimpleNamespace(type=PP_PLACEHOLDER.PICTURE)


class _GeometrylessShape:
    """A ``p:sp`` without geometry: python-pptx raises NotImplementedError on ``.shape_type``."""

    is_placeholder = False
    has_table = False
    top, left = Inches(2), Inches(1)
    has_text_frame = True
    text_frame = SimpleNamespace(paragraphs=[SimpleNamespace(level=0, text="Text in an odd shape")])

    @property
    def shape_type(self) -> object:
        raise NotImplementedError("Shape instance of unrecognized shape type")


def test_unreadable_shapes_are_skipped_and_the_slide_still_renders(tmp_path: Path) -> None:
    real = Presentation(str(_deck(tmp_path))).slides[0]
    shapes = _Shapes([*real.shapes, _LinkedPicture(), _LinkedPlaceholderPicture(), _GeometrylessShape()])
    shapes.title = real.shapes.title

    result = slide_to_markdown(SimpleNamespace(shapes=shapes), 1)

    assert result.markdown.startswith("## Graph basics")
    assert "- Vertices\n  - Edges connect | vertices" in result.markdown
    assert "Text in an odd shape" in result.markdown
    assert len(result.pictures) == 2  # the two embedded pictures; the linked ones are skipped


def test_normalize_picture_passthrough_convert_and_reject() -> None:
    png = png_bytes(20, 20)
    assert normalize_picture(png, "png") == (png, "png", "image/png")
    assert normalize_picture(b"\xff\xd8jpeg", "jpeg") == (b"\xff\xd8jpeg", "jpg", "image/jpeg")

    from PIL import Image

    gif = io.BytesIO()
    Image.new("P", (10, 10)).save(gif, "GIF")
    converted = normalize_picture(gif.getvalue(), "gif")
    assert converted is not None and converted[1] == "png" and converted[0].startswith(b"\x89PNG")

    assert normalize_picture(b"not an image (emf)", "x-emf") is None
